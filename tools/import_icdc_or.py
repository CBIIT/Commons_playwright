#!/usr/bin/env python3
"""
Refresh ICDC Object Repository for Commons_playwright.

Katalon still stores locators under Object Repository/Canine; Playwright
uses program name ICDC (lookup keys ICDC/...).

1. Import every Katalon .rs under Object Repository/Canine (CDS-style).
2. Overlay newer XPaths from commons-automation/ObjectRepository/ICDC YAML.
3. Soften study checkbox IDs so accession suffixes like (000015) still match.

Usage:
  python3 tools/import_icdc_or.py
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl
import yaml

REPO_ROOT = Path(__file__).resolve().parent.parent
PROGRAM = "ICDC"
KATALON_OR = REPO_ROOT.parent / "Commons_Automation" / "Object Repository" / "Canine"
YAML_OR = REPO_ROOT.parent / "commons-automation" / "ObjectRepository" / "ICDC"
OUT = REPO_ROOT / "ObjectRepository" / f"{PROGRAM}_ObjectRepository.xlsx"

# YAML group leaf → Katalon Filter folder
_YAML_FOLDER = {
    "Program": "Program",
    "Study": "Study",
    "StudyType": "StudyType",
    "Biobank": "Biobank",
    "StudyParticipation": "StudyParticipation",
    "Breed": "Breed",
    "Diagnosis": "Diagnosis",
    "PrimaryDiseaseSite": "PrimDiseaseSite",
    "StageOfDisease": "StageOfDisease",
    "ResponseToTreatment": "ResponseToTreatment",
    "Sex": "Sex",
    "NeuteredStatus": "NeuteredStatus",
    "SampleSite": "SampleSite",
    "SampleType": "SampleType",
    "SamplePathology": "SamplePathology",
    "FileAssociation": "FileAssociation",
    "FileType": "FileType",
    "FileFormat": "FileFormat",
}

_STUDY_ID_EXACT = re.compile(
    r"//\*\[@id=['\"]checkbox_Study_([^'\"]+)['\"]\]"
)


def _import_katalon() -> Path:
    sys.path.insert(0, str(REPO_ROOT / "tools"))
    from import_katalon_or import import_program

    return import_program(PROGRAM, KATALON_OR)


def _load_rows(path: Path) -> dict[str, list[tuple[str, str, str]]]:
    wb = openpyxl.load_workbook(path)
    by_sheet: dict[str, list[tuple[str, str, str]]] = {}
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[0] and row[2]:
                rows.append((str(row[0]).strip(), str(row[1] or "").strip(), str(row[2]).strip()))
        by_sheet[ws.title] = rows
    wb.close()
    return by_sheet


def _save_rows(path: Path, by_sheet: dict[str, list[tuple[str, str, str]]]) -> None:
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet_name in sorted(by_sheet.keys()):
        ws = wb.create_sheet(sheet_name[:31])
        ws.append(["Lookup Key", "Object Name", "XPath"])
        for row in sorted(by_sheet[sheet_name], key=lambda r: r[0]):
            ws.append(list(row))
        for col in range(1, 4):
            ws.column_dimensions[get_column_letter(col)].width = 40 if col == 1 else 28
    wb.save(path)


def _soften_study_xpath(xp: str) -> str:
    m = _STUDY_ID_EXACT.search(xp)
    if not m:
        return xp
    raw = m.group(1)
    code = re.sub(r"\s*\(\d+\)\s*$", "", raw).strip()
    if not code:
        return xp
    return f"//*[starts-with(@id, 'checkbox_Study_{code}')]"


def _index(by_sheet: dict[str, list[tuple[str, str, str]]]):
    by_lookup: dict[str, tuple[str, int]] = {}
    by_name: dict[str, list[tuple[str, int]]] = {}
    for sheet, rows in by_sheet.items():
        for i, (lk, name, _xp) in enumerate(rows):
            by_lookup[lk] = (sheet, i)
            by_name.setdefault(name, []).append((sheet, i))
            if name.endswith("_Chkbx"):
                by_name.setdefault(name[: -len("_Chkbx")], []).append((sheet, i))
    return by_lookup, by_name


def _set_xpath(by_sheet, sheet: str, idx: int, xpath: str) -> None:
    lk, name, _ = by_sheet[sheet][idx]
    by_sheet[sheet][idx] = (lk, name, xpath)


def _add_object(by_sheet, lookup: str, name: str, xpath: str) -> None:
    parts = lookup.split("/")
    sheet = parts[1] if len(parts) > 2 else "Root"
    by_sheet.setdefault(sheet, [])
    for i, (lk, _n, _x) in enumerate(by_sheet[sheet]):
        if lk == lookup:
            by_sheet[sheet][i] = (lookup, name, xpath)
            return
    by_sheet[sheet].append((lookup, name, xpath))


def _overlay_yaml(by_sheet) -> tuple[int, int]:
    if not YAML_OR.is_dir():
        print("[import_icdc_or] commons-automation ICDC YAML not found — Katalon only")
        return 0, 0
    by_lookup, by_name = _index(by_sheet)
    updated = 0
    added = 0

    def apply(lookup_candidates: list[str], name: str, xpath: str) -> None:
        nonlocal updated, added
        for lk in lookup_candidates:
            if lk in by_lookup:
                sheet, idx = by_lookup[lk]
                _set_xpath(by_sheet, sheet, idx, xpath)
                updated += 1
                return
        hits = by_name.get(name) or by_name.get(f"{name}_Chkbx")
        if hits:
            sheet, idx = hits[0]
            _set_xpath(by_sheet, sheet, idx, xpath)
            updated += 1
            return
        lk = lookup_candidates[0]
        _add_object(by_sheet, lk, name, xpath)
        added += 1
        by_lookup[lk] = (lk.split("/")[1] if lk.count("/") >= 2 else "Root", len(by_sheet.get(lk.split("/")[1], [])) - 1)

    for yml in sorted(YAML_OR.rglob("*.yaml")):
        data = yaml.safe_load(yml.read_text(encoding="utf-8")) or {}
        page = str(data.get("page") or "")
        for name, xp in (data.get("elements") or {}).items():
            if not xp:
                continue
            if page in ("(root)", "", "root"):
                apply([f"ICDC/{name}"], name, str(xp))
            else:
                apply([f"ICDC/{page}/{name}", f"ICDC/{name}"], name, str(xp))
        for group, gdata in (data.get("groups") or {}).items():
            leaf = str(group).split("/")[-1]
            folder = _YAML_FOLDER.get(leaf, leaf)
            if isinstance(gdata, dict):
                ddn = gdata.get("dropdown")
                if ddn:
                    apply(
                        [f"ICDC/Filter/{folder}/{leaf}_Ddn", f"ICDC/Filter/{folder}/Canine_Filter_{leaf}"],
                        f"{leaf}_Ddn",
                        str(ddn),
                    )
                for chk, xp in (gdata.get("checkboxes") or {}).items():
                    apply(
                        [
                            f"ICDC/Filter/{folder}/{chk}",
                            f"ICDC/Filter/{folder}/{chk}_Chkbx",
                        ],
                        f"{chk}_Chkbx" if not str(chk).endswith("_Chkbx") else str(chk),
                        str(xp),
                    )
                for extra, xp in (gdata.get("other") or {}).items():
                    apply([f"ICDC/{extra}", f"ICDC/Filter/{folder}/{extra}"], str(extra), str(xp))
    return updated, added


def _soften_all(by_sheet) -> int:
    n = 0
    for sheet, rows in by_sheet.items():
        for i, (lk, name, xp) in enumerate(rows):
            new_xp = _soften_study_xpath(xp)
            if new_xp != xp:
                by_sheet[sheet][i] = (lk, name, new_xp)
                n += 1
    return n


def main() -> None:
    if not KATALON_OR.is_dir():
        raise FileNotFoundError(KATALON_OR)
    _import_katalon()
    by_sheet = _load_rows(OUT)
    updated, added = _overlay_yaml(by_sheet)
    softened = _soften_all(by_sheet)
    # Accordion headers used when Sample/File facets sit under Filter By *
    _add_object(
        by_sheet,
        "ICDC/Filter/FilterByCases",
        "FilterByCases",
        "(//*[contains(normalize-space(.), 'Filter By Cases')])[1]",
    )
    _add_object(
        by_sheet,
        "ICDC/Filter/FilterBySamples",
        "FilterBySamples",
        "(//*[contains(normalize-space(.), 'Filter By Samples')])[1]",
    )
    _add_object(
        by_sheet,
        "ICDC/Filter/FilterByFiles",
        "FilterByFiles",
        "(//*[contains(normalize-space(.), 'Filter By Files')])[1]",
    )
    if "OSA04" not in " ".join(lk for rows in by_sheet.values() for lk, *_ in rows):
        _add_object(
            by_sheet,
            "ICDC/Filter/Study/Canine_Filter_Study-OSA04_Chkbx",
            "Canine_Filter_Study-OSA04_Chkbx",
            "//*[starts-with(@id, 'checkbox_Study_OSA04')]",
        )
        added += 1
    # Katalon groovy alias used by PRECINCT01 TC02
    _add_object(
        by_sheet,
        "ICDC/Filter/Study/PRECINCT01_Chkbx",
        "PRECINCT01_Chkbx",
        "//*[starts-with(@id, 'checkbox_Study_PRECINCT01')]",
    )
    _save_rows(OUT, by_sheet)
    total = sum(len(v) for v in by_sheet.values())
    print(
        f"[import_icdc_or] {total} objects | YAML updated={updated} added={added} "
        f"study-xpath softened={softened} -> {OUT}"
    )


if __name__ == "__main__":
    main()
