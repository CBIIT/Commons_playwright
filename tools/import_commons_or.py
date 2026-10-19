#!/usr/bin/env python3
"""
Import Object Repository YAML (commons-automation layout) into Playwright Excel OR.

Produces Katalon-compatible lookup keys so existing Excel test steps keep working:
  C3DC/Filters/Study/dbGaP_Accession/phs000463-Chkbx

Usage:
  python tools/import_commons_or.py C3DC \\
    ../commons-automation/ObjectRepository/C3DC

Re-run after commons-automation OR updates. Preserves legacy dropdown object names
(from the previous Excel file) and merges any entries missing from YAML.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl
import yaml
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = REPO_ROOT / "ObjectRepository"

CHKBX_SUFFIX = re.compile(r"[-_]Chkbx$")
DDN_SUFFIX = re.compile(r"[-_]Ddn$")

# Katalon facet element names on Filters/<section>/ (typo on TreatmentResp preserved).
_FACET_OBJ = {
    "Filters/Study": "Study_Facet",
    "Filters/Demographics": "Demographics_Facet",
    "Filters/Samples": "Samples_Facet",
    "Filters/Diagnosis": "Diagnosis-Facet",
    "Filters/Survival": "Survival-Facet",
    "Filters/Treatment": "Treatment-Facet",
    "Filters/TreatmentResp": "TreatReponse-Facet",
}


def _sheet_for_lookup(lookup: str) -> str:
    """Program / <sheet> / … — e.g. C3DC/Navbar/Explore-Tab → Navbar."""
    parts = lookup.split("/")
    if len(parts) < 3:
        return "Root"
    if parts[1] == "Static":
        return "Static"
    return parts[1]


def _load_legacy_ddn_names(xlsx: Path) -> dict[str, str]:
    """group_key (e.g. Study/dbGaP_Accession) -> dropdown object name (e.g. dbGaP_Accession-Ddn)."""
    if not xlsx.is_file():
        return {}
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    out: dict[str, str] = {}
    for ws in wb.worksheets:
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        try:
            lk_i, name_i = headers.index("Lookup Key"), headers.index("Object Name")
        except ValueError:
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            lk, obj = row[lk_i], row[name_i]
            if not lk or not obj or not DDN_SUFFIX.search(str(obj)):
                continue
            parts = str(lk).split("/")
            if len(parts) >= 5 and parts[1] == "Filters":
                group_key = f"{parts[2]}/{parts[3]}"
                out[group_key] = str(obj).strip()
    wb.close()
    return out


def _load_legacy_rows(xlsx: Path) -> dict[str, tuple[str, str, str]]:
    if not xlsx.is_file():
        return {}
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    rows: dict[str, tuple[str, str, str]] = {}
    for ws in wb.worksheets:
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        try:
            lk_i, name_i, xp_i = headers.index("Lookup Key"), headers.index("Object Name"), headers.index("XPath")
        except ValueError:
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            lk, name, xp = row[lk_i], row[name_i], row[xp_i]
            if lk and xp:
                rows[str(lk).strip()] = (str(lk).strip(), str(name or "").strip(), str(xp).strip())
    wb.close()
    return rows


def _ddn_object_name(group_key: str, legacy: dict[str, str]) -> str:
    if group_key in legacy:
        return legacy[group_key]
    folder = group_key.split("/")[-1]
    return f"{folder}-Ddn"


def _add(rows: dict, lookup: str, obj_name: str, xpath: str):
    xpath = " ".join(str(xpath).split())
    if lookup and xpath:
        rows[lookup] = (lookup, obj_name, xpath)


def _import_yaml_tree(program: str, yaml_root: Path, legacy_ddn: dict[str, str]) -> dict[str, tuple[str, str, str]]:
    rows: dict[str, tuple[str, str, str]] = {}
    facets_done: set[str] = set()

    for ypath in sorted(yaml_root.rglob("*.yaml")):
        doc = yaml.safe_load(ypath.read_text(encoding="utf-8")) or {}
        page = doc.get("page") or "(root)"
        if page == "(root)":
            prefix = ""
        else:
            prefix = page + "/"

        for name, xp in sorted((doc.get("elements") or {}).items()):
            lookup = f"{program}/{prefix}{name}"
            _add(rows, lookup, name, xp)

        page_elements = doc.get("elements") or {}
        for group_key, g in sorted((doc.get("groups") or {}).items()):
            folder = group_key.split("/")[-1]
            base = f"{program}/{prefix}{folder}"

            facet = g.get("facet")
            if not facet and g.get("facet_key"):
                facet = page_elements.get(g["facet_key"]) or rows.get(f"{program}/{prefix}{g['facet_key']}", (None, None, None))[2]
            if facet:
                top = "/".join(page.split("/")[:2]) if page.startswith("Filters/") else page
                facet_obj = _FACET_OBJ.get(top)
                if facet_obj and top not in facets_done:
                    facet_lookup = f"{program}/{top}/{facet_obj}"
                    _add(rows, facet_lookup, facet_obj, facet)
                    facets_done.add(top)

            if g.get("dropdown"):
                ddn_name = _ddn_object_name(group_key, legacy_ddn)
                _add(rows, f"{base}/{ddn_name}", ddn_name, g["dropdown"])
            for i, dd in enumerate(g.get("dropdowns") or []):
                name = f"{folder}-Ddn-{i + 1}" if i else f"{folder}-Ddn"
                _add(rows, f"{base}/{name}", name, dd)
            if g.get("search"):
                search_name = f"{folder}_Search"
                _add(rows, f"{base}/{search_name}", search_name, g["search"])

            for token, xp in sorted((g.get("checkboxes") or {}).items()):
                chk_name = f"{token}-Chkbx"
                _add(rows, f"{base}/{chk_name}", chk_name, xp)

            tpl = g.get("checkbox_template")
            if tpl:
                for token in g.get("values") or []:
                    chk_name = f"{token}-Chkbx"
                    _add(rows, f"{base}/{chk_name}", chk_name, tpl.replace("{value}", token))

            for oname, oxp in sorted((g.get("other") or {}).items()):
                _add(rows, f"{base}/{oname}", oname, oxp)

    return rows


def _write_xlsx(program: str, rows: dict[str, tuple[str, str, str]]) -> Path:
    by_sheet: dict[str, list[tuple[str, str, str]]] = {}
    for lookup, triple in rows.items():
        sheet = _sheet_for_lookup(lookup)
        by_sheet.setdefault(sheet, []).append(triple)

    out_path = OUT_DIR / f"{program}_ObjectRepository.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet_name in sorted(by_sheet.keys()):
        ws = wb.create_sheet(sheet_name[:31])
        ws.append(["Lookup Key", "Object Name", "XPath"])
        for triple in sorted(by_sheet[sheet_name], key=lambda r: r[0]):
            ws.append(list(triple))
        for col in range(1, 4):
            ws.column_dimensions[get_column_letter(col)].width = 44 if col == 1 else 28

    wb.save(out_path)
    return out_path


def import_program(program: str, yaml_root: Path) -> Path:
    yaml_root = Path(yaml_root)
    if not yaml_root.is_dir():
        raise FileNotFoundError(yaml_root)

    out_path = OUT_DIR / f"{program}_ObjectRepository.xlsx"
    legacy_ddn = _load_legacy_ddn_names(out_path)
    legacy_rows = _load_legacy_rows(out_path)

    rows = _import_yaml_tree(program, yaml_root, legacy_ddn)

    merged = 0
    for lk, triple in legacy_rows.items():
        if lk not in rows:
            rows[lk] = triple
            merged += 1

    out_path = _write_xlsx(program, rows)
    print(
        f"[import_commons_or] {program}: {len(rows)} objects -> {out_path}"
        + (f"  (+{merged} legacy-only entries kept)" if merged else "")
    )
    return out_path


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python tools/import_commons_or.py <Program> <commons-automation_OR_app_dir>")
        sys.exit(1)
    import_program(sys.argv[1], Path(sys.argv[2]))
