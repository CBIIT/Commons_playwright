#!/usr/bin/env python3
"""
Import Katalon Object Repository (.rs) trees into Playwright Excel OR files.

Usage:
  python tools/import_katalon_or.py C3DC \\
    "/Users/sohilz2/Automation/Commons_Automation/Object Repository/C3DC"

For C3DC, prefer commons-automation YAML (newer UI locators):
  python tools/import_commons_or.py C3DC ../commons-automation/ObjectRepository/C3DC
"""
import html
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = REPO_ROOT / "ObjectRepository"


def _xpath_from_rs(path: Path) -> str | None:
    text = path.read_text(encoding="utf-8", errors="replace")
    for pat in (
        r"<name>xpath</name>\s*<type>Main</type>\s*<value>(.*?)</value>",
        r"<key>BASIC</key>\s*<value>(.*?)</value>",
    ):
        m = re.search(pat, text, re.DOTALL | re.IGNORECASE)
        if m:
            xp = html.unescape(m.group(1).strip())
            xp = xp.replace("&quot;", '"').replace("&apos;", "'")
            return xp
    return None


def import_program(program: str, katalon_root: Path) -> Path:
    katalon_root = Path(katalon_root)
    if not katalon_root.is_dir():
        raise FileNotFoundError(katalon_root)

    by_sheet: dict[str, list[tuple[str, str, str]]] = {}

    for rs in sorted(katalon_root.rglob("*.rs")):
        rel = rs.relative_to(katalon_root)
        parts = rel.parts
        if not parts:
            continue
        sheet = parts[0] if len(parts) > 1 else "Root"
        obj_name = rs.stem
        subpath = "/".join(parts[:-1])
        lookup = f"{program}/{subpath}/{obj_name}" if subpath else f"{program}/{obj_name}"
        xpath = _xpath_from_rs(rs)
        if not xpath:
            continue
        by_sheet.setdefault(sheet, []).append((lookup, obj_name, xpath))

    out_path = OUT_DIR / f"{program}_ObjectRepository.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet_name in sorted(by_sheet.keys()):
        ws = wb.create_sheet(sheet_name[:31])
        ws.append(["Lookup Key", "Object Name", "XPath"])
        for row in sorted(by_sheet[sheet_name], key=lambda r: r[0]):
            ws.append(list(row))
        for col in range(1, 4):
            ws.column_dimensions[get_column_letter(col)].width = 40 if col == 1 else 28

    wb.save(out_path)
    total = sum(len(v) for v in by_sheet.values())
    print(f"[import_katalon_or] {program}: {total} objects -> {out_path}")
    return out_path


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python tools/import_katalon_or.py <Program> <Katalon_OR_root>")
        sys.exit(1)
    import_program(sys.argv[1], Path(sys.argv[2]))
