"""
excel_reader.py
Read input Excel files (study config) and compare WebData vs TsvData output sheets.
Replaces ReadExcel.groovy (initialLoad) and Utils.compareSheets() / compareTwoLists().

Input Excel columns (unchanged from Katalon):
  TabName  | TabQuery | StatQuery | TsvExcel | WebExcel
"""
import os
import re
import openpyxl
from config.settings import INPUT_FILES_ROOT, OUTPUT_FILES_ROOT


def load_study_config(input_excel_path: str) -> dict:
    """
    Read the input Excel and return all config the test needs.
    Replaces: ReadExcel.initialLoad() + excelparsingKatalon() in Katalon.

    Args:
        input_excel_path: Full path to input Excel e.g. InputFiles/CDS/TC01_CDS_phs001234_*.xlsx

    Returns:
        dict: {
            'input_excel':  str,   full path to this file
            'tsv_excel':    str,   full path to TsvData output Excel
            'web_excel':    str,   full path to WebData output Excel
            'tabs': {
                'ParticipantsTab': { 'tab_query': str, 'stat_query': str },
                'FilesTab':        { 'tab_query': str, 'stat_query': str },
                ...
            }
        }
    """
    wb = openpyxl.load_workbook(input_excel_path, data_only=True)
    ws = wb.active

    # Read header row to find column positions
    headers = [str(c.value).strip() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {h: i for i, h in enumerate(headers)}

    tabs       = {}
    tsv_excel  = ""
    web_excel  = ""

    for row in ws.iter_rows(min_row=2, values_only=True):
        tab_name = row[col.get("TabName", 0)]
        if not tab_name:
            continue

        tab_query  = row[col.get("TabQuery",  1)] or ""
        stat_query = row[col.get("StatQuery", 2)] or ""
        tsv_file   = row[col.get("TsvExcel",  3)] or ""
        web_file   = row[col.get("WebExcel",  4)] or ""

        # TsvExcel and WebExcel only appear on the first data row
        if tsv_file and not tsv_excel:
            tsv_excel = os.path.join(OUTPUT_FILES_ROOT, tsv_file)
        if web_file and not web_excel:
            web_excel = os.path.join(OUTPUT_FILES_ROOT, web_file)

        tabs[str(tab_name).strip()] = {
            "tab_query":  str(tab_query).strip(),
            "stat_query": str(stat_query).strip(),
        }

    wb.close()
    return {
        "input_excel": input_excel_path,
        "tsv_excel":   tsv_excel,
        "web_excel":   web_excel,
        "tabs":        tabs,
    }


def read_excel_sheet(excel_path: str, sheet_name: str) -> list:
    """
    Read an Excel sheet into a list of row lists (all values as strings).
    Replaces: ReadExcel.readOutputExcel() in Katalon.

    Args:
        excel_path: Full path to Excel file
        sheet_name: Sheet name to read

    Returns:
        list: [ ['col1', 'col2', ...], ['val1', 'val2', ...], ... ]
              First row is the header row.
    """
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    rows = []
    for row in ws.iter_rows(values_only=True):
        if any(v is not None for v in row):
            rows.append([str(v).strip() if v is not None else "" for v in row])
    wb.close()
    return rows


def compare_sheets(output_excel_path: str, web_sheet: str, tsv_sheet: str) -> dict:
    """
    Read WebData and TsvData sheets, sort both, compare row by row.
    Replaces: Utils.compareSheets() in Katalon.

    Args:
        output_excel_path: Full path to output Excel file
        web_sheet:         Name of the WebData sheet
        tsv_sheet:         Name of the TsvData sheet

    Returns:
        dict: {
            'passed':     bool,
            'mismatches': [ { 'row': int, 'col': str, 'web': str, 'tsv': str } ],
            'web_count':  int,
            'tsv_count':  int,
        }
    """
    web_rows = read_excel_sheet(output_excel_path, web_sheet)
    tsv_rows = read_excel_sheet(output_excel_path, tsv_sheet)

    # First row is headers — pull out and skip
    web_headers = web_rows[0] if web_rows else []
    web_data    = sorted([_normalize_row(r) for r in web_rows[1:]])
    tsv_data    = sorted([_normalize_row(r) for r in tsv_rows[1:]])

    mismatches = compare_two_lists(web_data, tsv_data, web_headers)

    return {
        "passed":     len(mismatches) == 0 and len(web_data) == len(tsv_data),
        "mismatches": mismatches,
        "web_count":  len(web_data),
        "tsv_count":  len(tsv_data),
    }


def compare_two_lists(web_rows: list, tsv_rows: list, headers: list = None) -> list:
    """
    Row-by-row comparison of web data vs TSV data. Returns list of mismatches.
    Replaces: Utils.compareTwoLists() in Katalon.

    Args:
        web_rows: Sorted rows from WebData sheet
        tsv_rows: Sorted rows from TsvData sheet
        headers:  Column header names for mismatch reporting

    Returns:
        list: [ { 'row': int, 'col': str, 'web': str, 'tsv': str } ]
    """
    mismatches = []
    max_rows   = max(len(web_rows), len(tsv_rows))

    for i in range(max_rows):
        web_row = web_rows[i] if i < len(web_rows) else []
        tsv_row = tsv_rows[i] if i < len(tsv_rows) else []

        if not web_row:
            mismatches.append({"row": i + 1, "col": "*", "web": "(missing row)", "tsv": str(tsv_row)})
            continue
        if not tsv_row:
            mismatches.append({"row": i + 1, "col": "*", "web": str(web_row), "tsv": "(missing row)"})
            continue

        max_cols = max(len(web_row), len(tsv_row))
        for j in range(max_cols):
            web_val = web_row[j] if j < len(web_row) else ""
            tsv_val = tsv_row[j] if j < len(tsv_row) else ""
            col_name = headers[j] if headers and j < len(headers) else f"Col{j+1}"

            web_clean = strip_cpi_badge(normalize_semicolon_list(web_val))
            tsv_clean = normalize_semicolon_list(tsv_val)

            if web_clean != tsv_clean:
                mismatches.append({
                    "row": i + 1,
                    "col": col_name,
                    "web": web_val,
                    "tsv": tsv_val,
                })

    return mismatches


def normalize_semicolon_list(value: str) -> str:
    """
    Sort semicolon-separated multi-values for order-independent comparison.
    e.g. 'B;A;C' → 'A;B;C'
    Replaces: Utils.normalizeSemicolonList() in Katalon.
    """
    if not value or ";" not in value:
        return value.strip() if value else ""
    parts = [p.strip() for p in value.split(";") if p.strip()]
    return ";".join(sorted(parts))


def strip_cpi_badge(value: str) -> str:
    """
    Remove CPI badge count suffix from UI cell values before comparison.
    e.g. 'StudyName [CPI: 3]' → 'StudyName'
    Replaces: Utils.stripCpiBadge() in Katalon.
    """
    if not value:
        return ""
    return re.sub(r'\s*\[CPI:\s*\d+\]', '', value).strip()


# ── Internal helpers ──────────────────────────────────────────────────────────

def _normalize_row(row: list) -> list:
    """Normalize every cell in a row for comparison."""
    return [strip_cpi_badge(normalize_semicolon_list(str(v))) for v in row]
