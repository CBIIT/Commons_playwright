"""
object_repo.py
Loads per-program Object Repository Excel files from ObjectRepository/ and
provides find_test_object() — the Playwright equivalent of Katalon's findTestObject().

Each program has its own file:  ObjectRepository/<Program>_ObjectRepository.xlsx
Each file has one sheet per page.  Each sheet has columns:
    Lookup Key | Object Name | XPath

Katalon:   findTestObject('CDS/NavBar/CDS_Data-Btn')
New:       find_test_object('CDS/NavBar/CDS_Data-Btn')  → returns XPath string
"""
import os
import openpyxl

# Root folder containing <Program>_ObjectRepository.xlsx files
OBJECT_REPO_DIR = os.path.join(
    os.path.dirname(os.path.dirname(__file__)),
    "ObjectRepository"
)

# { 'CDS': {'CDS/NavBar/CDS_Data-Btn': '//xpath', ...}, 'Canine': {...}, ... }
_repo: dict = {}


def _load_program(program: str) -> dict:
    """Load one program's Excel file — all sheets — into a flat lookup dict."""
    path = os.path.join(OBJECT_REPO_DIR, f"{program}_ObjectRepository.xlsx")
    if not os.path.exists(path):
        raise FileNotFoundError(
            f"No Object Repository found for program '{program}': {path}"
        )
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    objects = {}
    for ws in wb.worksheets:
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        try:
            lk_col = headers.index("Lookup Key")
            xp_col = headers.index("XPath")
        except ValueError:
            continue  # sheet has no expected columns — skip
        for row in ws.iter_rows(min_row=2, values_only=True):
            lk = row[lk_col]
            xp = row[xp_col]
            if lk and xp:
                objects[str(lk).strip()] = str(xp).strip()
    wb.close()
    return objects


def load_object_repo(program: str = None) -> dict:
    """
    Load Object Repository for one program (or all available programs if None).
    Results are cached — safe to call multiple times.

    Args:
        program: e.g. 'CDS', 'Canine', 'CTDC', or None to load all found files.

    Returns:
        dict: { lookup_key: xpath, ... } for the requested scope.
    """
    global _repo
    if program:
        if program not in _repo:
            _repo[program] = _load_program(program)
        return _repo[program]
    else:
        # Load every <Program>_ObjectRepository.xlsx found in the directory
        for fname in os.listdir(OBJECT_REPO_DIR):
            if fname.endswith("_ObjectRepository.xlsx"):
                prog = fname.replace("_ObjectRepository.xlsx", "")
                if prog not in _repo:
                    try:
                        _repo[prog] = _load_program(prog)
                        print(f"[object_repo] Loaded {prog}: {len(_repo[prog])} objects")
                    except Exception as e:
                        print(f"[object_repo] Warning: could not load {fname}: {e}")
        return {k: v for d in _repo.values() for k, v in d.items()}


def find_test_object(lookup_key: str, program: str = None) -> str:
    """
    Look up XPath by Katalon-style lookup key.
    Mirrors Katalon's findTestObject().

    Args:
        lookup_key: e.g. 'CDS/NavBar/CDS_Data-Btn'
        program:    Optional program name — skips search if already loaded.
                    Auto-derived from lookup_key prefix if not supplied.

    Returns:
        str: XPath expression

    Raises:
        KeyError:   lookup_key not found in the repository
        FileNotFoundError: program's Excel file does not exist
    """
    # Derive program from lookup key prefix if not given
    if not program:
        program = lookup_key.split("/")[0]

    # Lazy-load this program if not yet in cache
    if program not in _repo:
        _repo[program] = _load_program(program)

    if lookup_key not in _repo[program]:
        raise KeyError(
            f"Object not found: '{lookup_key}'\n"
            f"Check sheet(s) in ObjectRepository/{program}_ObjectRepository.xlsx"
        )
    return _repo[program][lookup_key]


def get_all_objects_for_program(program: str) -> dict:
    """Return all objects for a program as { lookup_key: xpath }."""
    if program not in _repo:
        _repo[program] = _load_program(program)
    return dict(_repo[program])


def reload_object_repo(program: str = None):
    """Force reload — evicts cache for one program (or all if None)."""
    global _repo
    if program:
        _repo.pop(program, None)
    else:
        _repo.clear()
    load_object_repo(program)


# ── Auto-load all available programs on import ────────────────────────────────
try:
    load_object_repo()
except Exception as e:
    print(f"[object_repo] Warning: auto-load failed: {e}")
