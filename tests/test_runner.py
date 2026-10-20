"""
test_runner.py
Excel-driven test orchestrator. Reads TC Excel Steps sheet and executes each action.

Steps sheet columns:
  # | Program | Input Excel | Max Pages | Max Rows | Action | Page | Object Name | Params | Filter | Notes

Filter column: tester writes key=value on checkbox steps e.g. sex=Male, phs=phs001437
Runner collects all Filter values during execution, then uses them to:
  - Build Cypher via core/query_builders.py (per-app: query_builder_cds.py, query_builder_c3dc.py, …)
  - Query Memgraph via memgraph_reader.py
  - Compare UI rows vs Memgraph rows on multi_function steps
  - Compare UI stat bar counts vs Memgraph StatBar query on read_stat_bar steps

If a Queries sheet exists in the TC Excel, its Cypher is used instead of the template (one-off override).
"""
import os
import openpyxl
from core.object_repo   import find_test_object
from core.browser       import launch_browser, navigate_to, close_browser
from core.actions       import dismiss_popup, click_element, click_and_wait
from core.stat_bar      import read_stat_bar, validate_stat_bar
from core.table_reader  import collect_all_rows, write_web_data
from core.excel_reader  import load_study_config, compare_sheets
from config.settings    import get_url, OUTPUT_FILES_ROOT

# ── Per-program result tab definitions ────────────────────────────────────────
def _c3dc_tab(tab_name, tab_label, tbl_key, hdr_key, stat_key=None):
    """Helper: C3DC explore tabs use table OR keys + shared Next button + tab_label for More menu."""
    return {
        "tab_name":   tab_name,
        "tab_label":  tab_label,
        "tab_btn":    None,
        "table_body": f"C3DC/ResultTabs/{tbl_key}",
        "table_hdr":  f"C3DC/ResultTabs/{hdr_key}",
        "next_btn":   "C3DC/ResultTabs/All_Tabs_Next-Btn",
        "stat_key":   stat_key,
    }


PROGRAM_TABS = {
    "CDS": [
        {
            "tab_name":   "ParticipantsTab",
            "tab_btn":    "CDS/Data_page/CDSResults_Participants_Tab",
            "table_body": "CDS/Data_page/CDS_ParticipantsTable",
            "table_hdr":  "CDS/Data_page/CDS_ParticipantsTableHeader",
            "next_btn":   "CDS/Data_page/CDS_ParticipantsTabNextBtn",
        },
        {
            "tab_name":   "SamplesTab",
            "tab_btn":    "CDS/Data_page/CDSResults_Samples_Tab",
            "table_body": "CDS/Data_page/CDS_SamplesTable",
            "table_hdr":  "CDS/Data_page/CDS_SamplesTableHeader",
            "next_btn":   "CDS/Data_page/CDS_SamplesTabNextBtn",
        },
        {
            "tab_name":   "FilesTab",
            "tab_btn":    "CDS/Data_page/CDSResults_Files_Tab",
            "table_body": "CDS/Data_page/CDS_FilesTable",
            "table_hdr":  "CDS/Data_page/CDS_FilesTableHeader",
            "next_btn":   "CDS/Data_page/CDS_FilesTabNextBtn",
        },
    ],
    "C3DC": [
        _c3dc_tab("ParticipantsTab", "Participants", "Participants-Tbl", "Participants-TblHdr", stat_key="participants"),
        _c3dc_tab("StudiesTab", "Studies", "Studies-Tbl", "Studies-TblHdr", stat_key="studies"),
        _c3dc_tab("DiagnosisTab", "Diagnosis", "Diagnosis-Tbl", "Diagnosis-TblHdr", stat_key="participants"),
        _c3dc_tab("GeneticAnalysisTab", "Genetic Analyses", "GeneticAnalysis-Tbl", "GeneticAnalysis-TblHdr", stat_key="participants"),
        _c3dc_tab("TreatmentTab", "Treatments", "Treatment-Tbl", "Treatment-TblHdr", stat_key="participants"),
        _c3dc_tab("TreatmentRespTab", "Treatment Response", "TreatmentResp-Tbl", "TreatmentResp-TblHdr", stat_key="participants"),
        _c3dc_tab("SurvivalTab", "Survival", "Survival-Tbl", "Survival-TblHdr", stat_key="participants"),
        _c3dc_tab("SamplesTab", "Samples", "Samples-Tbl", "Samples-TblHdr", stat_key="participants"),
        _c3dc_tab("FilesTab", "Files", "Files-Tbl", "Files-TblHdr", stat_key="files"),
    ],
}

PROGRAM_NAV_BTN = {
    "CDS":    "CDS/NavBar/CDS_Data-Btn",
    "Canine": "Canine/NavBar/Canine_Data-Btn",
    "CTDC":   "CTDC/NavBar/CTDC_Data-Btn",
    "C3DC":   "C3DC/Navbar/Explore-Tab",
}


def run_test_from_excel(testcase_excel: str, run_dir: str = None):
    """
    Read a TC Excel (Steps sheet) and execute every step.
    Main entry point for single TC runs.

    Args:
        testcase_excel: path to TC Excel e.g. TestCases/TC01_CDS_phs001437.xlsx
        run_dir:        optional folder to write the output Excel into (e.g. a
                        timestamped Reports/run_YYYYMMDD_HHMMSS/ subfolder).
                        Defaults to Reports/ when None.
    """
    wb = openpyxl.load_workbook(testcase_excel, data_only=True)
    sheet_names = wb.sheetnames

    # ── Read Steps sheet ──────────────────────────────────────────────────────
    # Row 1 = column headers (ignored at runtime)
    # Row 2 = config values: # | Program | Input Excel | Max Pages | Max Rows | (rest empty)
    # Row 3+ = steps: # | (empty) | (empty) | (empty) | (empty) | Action | Page | Object Name | Params | Filter | Notes
    ws = wb["Steps"]
    rows = list(ws.iter_rows(min_row=1, values_only=True))

    cfg_row   = rows[1]   # row 2 in the Excel (index 1)
    program   = str(cfg_row[1]).strip()  if cfg_row[1] else ""
    input_excel = str(cfg_row[2]).strip() if cfg_row[2] else ""
    max_pages = int(cfg_row[3]) if cfg_row[3] else None
    max_rows  = int(cfg_row[4]) if cfg_row[4] else None

    steps = []
    for row in rows[2:]:
        if not row[5]:
            continue
        steps.append({
            "action":      str(row[5]).strip(),
            "page":        str(row[6]).strip() if row[6] else "",
            "object_name": str(row[7]).strip() if row[7] else "",
            "params":      str(row[8]).strip() if row[8] else "",
            "filter":      str(row[9]).strip() if row[9] else "",
            "notes":       str(row[10]).strip() if len(row) > 10 and row[10] else "",
        })

    # ── Read Queries sheet if present (one-off Cypher override) ──────────────
    # Format: Tab | Cypher Query
    custom_queries = {}
    if "Queries" in sheet_names:
        ws_q = wb["Queries"]
        for qrow in ws_q.iter_rows(min_row=2, values_only=True):
            if qrow[0] and qrow[1]:
                custom_queries[str(qrow[0]).strip()] = str(qrow[1]).strip()
        print(f"[test_runner] Queries sheet found — custom queries for: {list(custom_queries.keys())}")

    wb.close()

    from config.settings import ENV
    url = get_url(program)
    print(f"\n{'='*60}")
    print(f"[test_runner] TC: {os.path.basename(testcase_excel)}")
    print(f"[test_runner] Program={program}  ENV={ENV}  URL={url}")
    print(f"[test_runner] Steps={len(steps)}  Custom queries={len(custom_queries)}")
    print(f"{'='*60}")

    # Resolve input excel path
    from config.settings import COMMONS_ROOT
    if input_excel and not os.path.isabs(input_excel):
        input_excel = os.path.join(COMMONS_ROOT, input_excel)

    playwright, browser, page = None, None, None
    stat_counts  = {}
    collected_filters: dict = {}   # {key: [value, ...]} — accumulated from Filter column
    result = {"passed": True, "tabs": {}, "errors": [], "stat_bar": []}

    try:
        playwright, browser, page = launch_browser()
        navigate_to(page, url)

        for step in steps:
            action   = step["action"]
            page_name = step["page"]
            obj_name  = step["object_name"]
            params    = step["params"]
            flt       = step["filter"]

            def _xpath():
                if not page_name or not obj_name:
                    raise ValueError(f"Step '{action}' needs Page + Object Name but one is empty.")
                # Strip leading page_name prefix from obj_name if present (avoids doubling)
                # e.g. Page="Data_page", ObjName="Data_page/Filter/..." → "Filter/..."
                effective_obj = obj_name
                if effective_obj.startswith(page_name + "/"):
                    effective_obj = effective_obj[len(page_name) + 1:]
                return find_test_object(f"{program}/{page_name}/{effective_obj}")

            print(f"[test_runner] {action}"
                  f"{f' | {page_name}/{obj_name}' if obj_name else ''}"
                  f"{f' | params={params}' if params else ''}"
                  f"{f' | filter={flt}' if flt else ''}")

            # ── Collect filter value from this step ───────────────────────────
            if flt and "=" in flt:
                key, val = flt.split("=", 1)
                key = key.strip()
                val = val.strip()
                collected_filters.setdefault(key, [])
                if val not in collected_filters[key]:
                    collected_filters[key].append(val)

            # ── Execute action ────────────────────────────────────────────────
            if action == "dismiss_popup":
                dismiss_popup(page)

            elif action == "navigate_to":
                from config.settings import get_page_url
                navigate_to(page, get_page_url(program, page_name))

            elif action == "click_element":
                _ensure_facet_expanded(page, program, page_name, obj_name)
                click_element(page, _xpath())

            elif action == "click_and_wait":
                _ensure_facet_expanded(page, program, page_name, obj_name)
                click_and_wait(page, _xpath())
                if program == "C3DC" and page_name == "Navbar" and obj_name == "Explore-Tab":
                    _wait_c3dc_explore_ready(page)

            elif action == "js_click_and_wait":
                from core.actions import js_click
                _ensure_facet_expanded(page, program, page_name, obj_name)
                js_click(page, _xpath())
                page.wait_for_load_state("networkidle", timeout=15000)

            elif action == "js_click":
                from core.actions import js_click
                js_click(page, _xpath())

            elif action == "scroll_into_view":
                from core.actions import scroll_into_view
                scroll_into_view(page, _xpath())

            elif action == "wait_for_element_hidden":
                page.locator(f"xpath={_xpath()}").wait_for(state="hidden")

            elif action in ("is_element_visible", "is_element_enabled"):
                loc = page.locator(f"xpath={_xpath()}")
                val = loc.is_enabled() if action == "is_element_enabled" else loc.is_visible()
                print(f"[test_runner]   => {val}")

            elif action == "clear_input":
                page.locator(f"xpath={_xpath()}").clear()

            elif action == "find_filter_by_search":
                from core.actions import find_filter_by_search
                find_filter_by_search(page, _xpath(), params)

            elif action == "read_stat_bar":
                # Read UI stat bar
                stat_counts = read_stat_bar(page, params or program)
                print(f"[test_runner]   UI stat bar: {stat_counts}")

                # Compare with Memgraph StatBar query if filters collected
                if collected_filters:
                    result["stat_bar"] = _compare_stat_bar(program, collected_filters, stat_counts, custom_queries)

            elif action == "multi_function":
                tab_name = params
                tab_def  = next((t for t in PROGRAM_TABS.get(program, [])
                                  if t["tab_name"] == tab_name), None)
                if not tab_def:
                    print(f"[test_runner] Warning: no tab definition for '{tab_name}' — skipping")
                    continue

                if program == "C3DC" and tab_name == "FilesTab":
                    from core.actions import click_and_wait
                    click_and_wait(page, find_test_object("C3DC/Navbar/ExploreFiles-Tab"))
                    page.wait_for_timeout(800)
                if program == "C3DC" and tab_def.get("tab_label"):
                    from core.c3dc_tabs import select_c3dc_tab
                    select_c3dc_tab(page, tab_def["tab_name"])
                else:
                    click_element(page, find_test_object(tab_def["tab_btn"]))

                tab_result = _run_multi_function(
                    page=page,
                    program=program,
                    tab_def=tab_def,
                    stat_counts=stat_counts,
                    filters=collected_filters,
                    custom_query=custom_queries.get(tab_name),
                    max_pages=max_pages,
                    max_rows=max_rows,
                )
                result["tabs"][tab_name] = tab_result
                if not tab_result["passed"]:
                    result["passed"] = False

            elif action == "login":
                parts = params.split(",")
                _login(page, parts[0].strip(), parts[1].strip() if len(parts) > 1 else "")

            elif action == "verify_static_page":
                ok = _verify_static(page, program, [params])
                if not ok:
                    result["passed"] = False

            else:
                print(f"[test_runner] Warning: unknown action '{action}' — skipping")

    except Exception as e:
        result["passed"] = False
        result["errors"].append(str(e))
        print(f"[test_runner] ERROR: {e}")
        raise
    finally:
        if browser and playwright:
            close_browser(playwright, browser)

    # ── Write output Excel with UI + DB data ─────────────────────────────────
    _write_output_excel(
        testcase_excel   = testcase_excel,
        program          = program,
        filters          = collected_filters,
        stat_counts      = stat_counts,
        result           = result,
        run_dir          = run_dir,
    )

    status = "PASS" if result["passed"] else "FAIL"
    print(f"\n{'='*60}")
    print(f"[test_runner] Result: {status}")
    for tab_name, tab_res in result["tabs"].items():
        tab_status = "PASS" if tab_res["passed"] else "FAIL"
        mm = len(tab_res.get("mismatches", []))
        print(f"  {tab_name}: {tab_status}{f'  ({mm} mismatches)' if mm else ''}")
    if result["errors"]:
        for e in result["errors"]:
            print(f"  ERROR: {e}")
    print(f"{'='*60}\n")
    return result


def _run_multi_function(page, program: str, tab_def: dict, stat_counts: dict,
                         filters: dict, custom_query: str = None,
                         max_pages: int = None, max_rows: int = None) -> dict:
    """
    Paginate UI result tab, query Memgraph with same filters, compare row counts.

    Returns:
        dict: { passed, mismatches, ui_count, db_count, ui_rows, db_rows, headers }
    """
    from core.query_builders import build_query
    from core.memgraph_reader import query_memgraph

    tab_name  = tab_def["tab_name"]
    stat_key  = tab_def.get("stat_key") or _tab_to_stat_key(tab_name)
    stat_value = stat_counts.get(stat_key, 0)

    if stat_value == 0:
        print(f"[test_runner]   Skipping '{tab_name}' — stat bar count is 0")
        return {"passed": True, "mismatches": [], "ui_count": 0, "db_count": 0,
                "ui_rows": [], "db_rows": [], "headers": [], "skipped": True}

    # ── Collect UI rows ───────────────────────────────────────────────────────
    ui_result = collect_all_rows(
        page             = page,
        table_body_xpath = find_test_object(tab_def["table_body"]) + "//tbody//tr",
        next_btn_xpath   = find_test_object(tab_def["next_btn"]),
        header_xpath     = find_test_object(tab_def["table_hdr"]) + "//th",
        expected_total   = stat_value,
        max_pages        = max_pages,
        max_rows         = max_rows,
    )
    ui_rows   = ui_result["rows"]
    headers   = ui_result.get("headers", [])
    truncated = ui_result.get("truncated", False)
    if truncated:
        ui_count = _read_pagination_total(page) or stat_value
        print(f"[test_runner]   UI rows truncated — pagination total: {ui_count}")
    else:
        ui_count = len(ui_rows)
        print(f"[test_runner]   UI rows collected: {ui_count}")

    # ── Query Memgraph (with retry on suspicious 0-row result) ───────────────
    import time as _time
    _RETRY_ATTEMPTS = 3
    _RETRY_DELAY    = 4  # seconds
    try:
        cypher = build_query(program, tab_name, filters, custom_query=custom_query)
        db_rows  = query_memgraph(cypher)
        db_count = len(db_rows)
        # If DB returned 0 but UI shows rows, Memgraph may have silently dropped
        # the query under concurrent load — retry before declaring a mismatch.
        if db_count == 0 and ui_count > 0:
            for _attempt in range(2, _RETRY_ATTEMPTS + 1):
                print(f"[test_runner]   DB=0 but UI={ui_count}, retrying Memgraph (attempt {_attempt}/{_RETRY_ATTEMPTS}) in {_RETRY_DELAY}s…")
                _time.sleep(_RETRY_DELAY)
                db_rows  = query_memgraph(cypher)
                db_count = len(db_rows)
                if db_count > 0:
                    break
        print(f"[test_runner]   Memgraph rows returned: {db_count}")
    except Exception as e:
        print(f"[test_runner]   Memgraph query failed: {e}")
        return {"passed": False, "mismatches": [str(e)], "ui_count": ui_count, "db_count": 0,
                "ui_rows": ui_rows, "db_rows": [], "headers": headers}

    # ── Compare row counts ────────────────────────────────────────────────────
    count_match = (ui_count == db_count)
    mismatches  = []
    if not count_match:
        mismatches.append(f"Row count mismatch: UI={ui_count} Memgraph={db_count}")

    passed = count_match
    status = "PASS" if passed else "FAIL"
    print(f"[test_runner]   {tab_name}: {status} | UI={ui_count} DB={db_count}")
    if mismatches:
        for m in mismatches:
            print(f"[test_runner]     MISMATCH: {m}")

    return {
        "passed":     passed,
        "mismatches": mismatches,
        "ui_count":   ui_count,
        "db_count":   db_count,
        "ui_rows":    ui_rows,
        "db_rows":    db_rows,
        "headers":    headers,
        "truncated":  truncated,
    }


def _compare_stat_bar(program: str, filters: dict, ui_counts: dict, custom_queries: dict) -> list[dict]:
    """
    Query Memgraph StatBar and compare counts against UI stat bar.
    Returns list of comparison rows: [{metric, ui, db, match}]
    """
    from core.query_builders import build_query
    from core.memgraph_reader import query_memgraph

    comparison = []
    try:
        custom = custom_queries.get("StatBar")
        cypher = build_query(program, "StatBar", filters, custom_query=custom)
        rows   = query_memgraph(cypher)
        if not rows:
            print("[test_runner]   StatBar: no rows returned from Memgraph")
            return comparison
        db = rows[0]
        print(f"[test_runner]   Memgraph stat bar: {dict(db)}")

        mapping = {
            "Studies":      "studies",
            "Participants": "participants",
            "Samples":      "samples",
            "Files":        "files",
        }
        for db_key, ui_key in mapping.items():
            db_val = db.get(db_key)
            ui_val = ui_counts.get(ui_key)
            if db_val is not None and ui_val is not None:
                matched = (db_val == ui_val)
                print(f"[test_runner]   {db_key}: UI={ui_val} DB={db_val} {'✓' if matched else '✗ MISMATCH'}")
                comparison.append({"metric": db_key, "ui": ui_val, "db": db_val, "match": matched})
    except Exception as e:
        print(f"[test_runner]   StatBar Memgraph query failed: {e}")
    return comparison


def _read_pagination_total(page) -> int:
    """
    Read the "of X" total from any pagination widget visible on the page.
    Looks for text like '1-10 of 159' or 'Showing 1 to 10 of 159'.
    Returns the total integer, or 0 if not found.
    """
    import re
    try:
        # Generic MUI / common pagination patterns
        locators = [
            "xpath=//*[contains(@class,'pagination') or contains(@class,'Pagination')]"
            "[contains(text(),' of ') or contains(text(),'of ')]",
            "xpath=//p[contains(text(),' of ')]",
            "xpath=//span[contains(text(),' of ')]",
        ]
        for xpath in locators:
            try:
                el = page.locator(xpath).first
                if el.is_visible():
                    text = el.inner_text().strip()
                    m = re.search(r'of\s+([\d,]+)', text, re.IGNORECASE)
                    if m:
                        return int(m.group(1).replace(",", ""))
            except Exception:
                continue
    except Exception:
        pass
    return 0


def _tab_to_stat_key(tab_name: str) -> str:
    mapping = {
        "ParticipantsTab": "participants",
        "SamplesTab":      "samples",
        "FilesTab":        "files",
        "StudiesTab":      "studies",
    }
    return mapping.get(tab_name, tab_name.lower().replace("tab", "").strip())


def _login(page, email: str, password: str):
    btn = page.locator("xpath=//button[contains(text(),'Sign In') or contains(text(),'Login')]").first
    btn.wait_for(state="visible")
    btn.click()
    page.locator("xpath=//input[@type='email' or @name='email']").first.fill(email)
    page.keyboard.press("Enter")
    page.locator("xpath=//input[@type='password']").first.fill(password)
    page.keyboard.press("Enter")
    page.wait_for_load_state("networkidle")


def _write_output_excel(testcase_excel: str, program: str, filters: dict,
                         stat_counts: dict, result: dict, run_dir: str = None):
    """
    Write a per-TC output Excel file to run_dir (or Reports/ if None) with:
      - Summary sheet  : TC info, filters, stat bar counts, per-tab pass/fail
      - UI_<Tab> sheet : rows collected from the browser (up to max_pages)
      - DB_<Tab> sheet : rows returned from Memgraph
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from datetime import datetime
    from pathlib import Path

    RED   = PatternFill("solid", fgColor="FFCCCC")
    GREEN = PatternFill("solid", fgColor="CCFFCC")
    BOLD  = Font(bold=True)

    if run_dir:
        reports_dir = Path(run_dir)
    else:
        reports_dir = Path(__file__).parent.parent / "Reports"
    reports_dir.mkdir(parents=True, exist_ok=True)
    tc_stem  = Path(testcase_excel).stem
    ts       = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = reports_dir / f"{tc_stem}_{ts}.xlsx"

    wb = openpyxl.Workbook()
    wb.remove(wb.active)   # remove default sheet

    # ── Summary sheet ─────────────────────────────────────────────────────────
    ws_sum = wb.create_sheet("Summary")
    overall_pass = result.get("passed", False)

    header_pairs = [
        ("TC",       tc_stem),
        ("Program",  program),
        ("Status",   "PASS" if overall_pass else "FAIL"),
        ("Run at",   datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("", ""),
        ("Filters applied", ""),
    ]
    for label, value in header_pairs:
        row = ws_sum.append([label, value]) or ws_sum.max_row
    for k, vals in filters.items():
        ws_sum.append(["", f"  {k} = {', '.join(vals)}"])

    ws_sum.append(["", ""])
    ws_sum.append(["Stat Bar Counts", ""])
    for k, v in stat_counts.items():
        ws_sum.append(["", f"  {k}: {v}"])

    ws_sum.append(["", ""])
    ws_sum.append(["Tab Results", "UI Count", "DB Count", "Status"])
    for tab_name, tab_res in result.get("tabs", {}).items():
        status_str = "PASS" if tab_res.get("passed") else "FAIL"
        r = ws_sum.max_row + 1
        ws_sum.append([tab_name, tab_res.get("ui_count", ""), tab_res.get("db_count", ""), status_str])
        fill = GREEN if tab_res.get("passed") else RED
        for col in range(1, 5):
            ws_sum.cell(r, col).fill = fill
        for mm in tab_res.get("mismatches", []):
            ws_sum.append(["", "", "", mm])

    # Bold first column
    for row in ws_sum.iter_rows():
        row[0].font = BOLD
    ws_sum.column_dimensions["A"].width = 28
    ws_sum.column_dimensions["B"].width = 20

    # ── Data sheets per tab ───────────────────────────────────────────────────
    for tab_name, tab_res in result.get("tabs", {}).items():
        if tab_res.get("skipped"):
            continue
        short = tab_name.replace("Tab", "")   # "Participants", "Samples", "Files"

        ui_rows  = tab_res.get("ui_rows", [])
        db_rows  = tab_res.get("db_rows", [])
        headers  = tab_res.get("headers", [])
        truncated = tab_res.get("truncated", False)

        # UI sheet
        ws_ui = wb.create_sheet(f"UI_{short}")
        note = f"(first {len(ui_rows)} of {tab_res.get('ui_count','?')} — truncated at max_pages)" if truncated else ""
        if note:
            ws_ui.append([note])
            ws_ui.cell(ws_ui.max_row, 1).font = Font(italic=True, color="888888")
        if ui_rows:
            if headers:
                ws_ui.append(headers)
                for cell in ws_ui[ws_ui.max_row]:
                    cell.font = BOLD
                    cell.fill = PatternFill("solid", fgColor="D9E1F2")
            for row in ui_rows:
                # ui_rows are lists when headers are separate, dicts otherwise
                ws_ui.append(list(row.values()) if isinstance(row, dict) else list(row))
        else:
            ws_ui.append(["(no rows collected)"])

        # DB sheet
        ws_db = wb.create_sheet(f"DB_{short}")
        if db_rows:
            db_cols = list(db_rows[0].keys())
            ws_db.append(db_cols)
            for cell in ws_db[1]:
                cell.font = BOLD
                cell.fill = PatternFill("solid", fgColor="E2EFDA")
            for row in db_rows:
                ws_db.append([row.get(c, "") for c in db_cols])
        else:
            ws_db.append(["(no rows returned)"])

    # ── Cell-by-cell comparison sheets ───────────────────────────────────────
    MISMATCH_FILL = PatternFill("solid", fgColor="FFD700")
    MATCH_FILL    = PatternFill("solid", fgColor="CCFFCC")
    for tab_name, tab_res in result.get("tabs", {}).items():
        if tab_res.get("skipped"):
            continue
        ui_rows = tab_res.get("ui_rows", [])
        db_rows = tab_res.get("db_rows", [])
        headers = tab_res.get("headers", [])
        if not ui_rows or not db_rows:
            continue

        short = tab_name.replace("Tab", "")
        ws_cmp = wb.create_sheet(f"Compare_{short}")

        # Build DB rows as lists in same column order as headers
        db_cols = list(db_rows[0].keys()) if db_rows else []
        # Map DB col index by position (headers align to DB cols after offset fix)
        col_count = max(len(headers), len(db_cols))

        # Header row
        cmp_headers = []
        for h in headers:
            cmp_headers += [f"UI_{h}", f"DB_{h}", "Match"]
        ws_cmp.append(cmp_headers)
        for cell in ws_cmp[1]:
            cell.font = BOLD
            cell.fill = PatternFill("solid", fgColor="D9E1F2")

        # Compare row by row (only collected UI rows)
        for i, ui_row in enumerate(ui_rows):
            db_row = db_rows[i] if i < len(db_rows) else {}
            db_vals = [db_row.get(c, "") for c in db_cols]
            out = []
            for j, ui_val in enumerate(ui_row):
                db_val = db_vals[j] if j < len(db_vals) else ""
                match = str(ui_val).strip() == str(db_val).strip()
                out += [ui_val, db_val, "✓" if match else "✗"]
            ws_cmp.append(out)
            # Colour mismatched triples
            for j in range(len(ui_row)):
                match_cell = ws_cmp.cell(ws_cmp.max_row, j * 3 + 3)
                if match_cell.value == "✗":
                    for k in range(3):
                        ws_cmp.cell(ws_cmp.max_row, j * 3 + 1 + k).fill = MISMATCH_FILL
                else:
                    ws_cmp.cell(ws_cmp.max_row, j * 3 + 3).fill = MATCH_FILL

    # ── StatBar comparison sheet ──────────────────────────────────────────────
    stat_bar_rows = result.get("stat_bar", [])
    if stat_bar_rows:
        ws_sb = wb.create_sheet("StatBar")
        header = ["Metric", "UI Value", "DB Value", "Match"]
        ws_sb.append(header)
        for cell in ws_sb[1]:
            cell.font = BOLD
            cell.fill = PatternFill("solid", fgColor="D9E1F2")
        for sb in stat_bar_rows:
            matched = sb.get("match", False)
            ws_sb.append([sb["metric"], sb["ui"], sb["db"], "PASS" if matched else "MISMATCH"])
            fill = GREEN if matched else RED
            for col in range(1, 5):
                ws_sb.cell(ws_sb.max_row, col).fill = fill
        ws_sb.column_dimensions["A"].width = 16
        ws_sb.column_dimensions["B"].width = 14
        ws_sb.column_dimensions["C"].width = 14
        ws_sb.column_dimensions["D"].width = 12

        # Flag stat bar mismatches on Summary sheet
        if any(not sb.get("match") for sb in stat_bar_rows):
            ws_sum.append(["", ""])
            r = ws_sum.max_row + 1
            ws_sum.append(["StatBar", "MISMATCH — see StatBar sheet"])
            for col in range(1, 3):
                ws_sum.cell(r, col).fill = RED

    wb.save(out_path)
    print(f"[test_runner] Output Excel: {out_path}")


# Maps facet folder name → Object Repo key for the section expander button.
# If the expander element is not visible (section already open), this is a no-op.
_FACET_EXPANDERS = {
    "CDS": {
        "DemographicsFacet": "CDS/Data_page/Filter/DemographicsFacet/DemographicsFacet",
        "GenomicFacet/Platform": "CDS/Data_page/Filter/GenomicFacet/Platform/platform_Ddn",
        "GenomicFacet/LibraryStrategy": "CDS/Data_page/Filter/GenomicFacet/LibraryStrategy/Library_Strategy_Ddn",
        "GenomicFacet/LibrarySelection": "CDS/Data_page/Filter/GenomicFacet/LibrarySelection/LibrarySelection_Ddn",
        "GenomicFacet/LibraryLayout": "CDS/Data_page/Filter/GenomicFacet/LibraryLayout/LibraryLayout_Ddn",
        "GenomicFacet/InstrumentModel": "CDS/Data_page/Filter/GenomicFacet/InstrumentModel/Instrument_model_Ddn",
        "GenomicFacet/ReferenceGenomeAssembly": "CDS/Data_page/Filter/GenomicFacet/ReferenceGenomeAssembly/Ref_Gen_Assembly_Ddn",
        "GenomicFacet/LibrarySource": "CDS/Data_page/Filter/GenomicFacet/LibrarySource/LibrarySourceMolecule_Ddn",
        "SamplesFacet/SampleTumorStatus": "CDS/Data_page/Filter/SamplesFacet/SampleTumorStatus/SampleTumorStatus-Ddn",
        "SamplesFacet/SampleType": "CDS/Data_page/Filter/SamplesFacet/SampleType/SampleType-Ddn",
        "FilesFacet/FileType": "CDS/Data_page/Filter/FilesFacet/FileType/FileType_Ddn",
        "SequencingFacet/Platform": "CDS/Data_page/Filter/SequencingFacet/Platform/platform_Ddn",
        "SequencingFacet/LibraryStrategy": "CDS/Data_page/Filter/SequencingFacet/LibraryStrategy/Library_Strategy_Ddn",
        "SequencingFacet/LibrarySelection": "CDS/Data_page/Filter/SequencingFacet/LibrarySelection/LibrarySelection_Ddn",
        "SequencingFacet/LibraryLayout": "CDS/Data_page/Filter/SequencingFacet/LibraryLayout/LibraryLayout_Ddn",
        "SequencingFacet/InstrumentModel": "CDS/Data_page/Filter/SequencingFacet/InstrumentModel/Instrument_model_Ddn",
        "SequencingFacet/ReferenceGenomeAssembly": "CDS/Data_page/Filter/SequencingFacet/ReferenceGenomeAssembly/Ref_Gen_Assembly_Ddn",
    },
    "C3DC": {
        "Study/dbGaP_Accession": "C3DC/Filters/Study/dbGaP_Accession/dbGaP_Accession-Ddn",
        "Demographics/SexAtBirth": "C3DC/Filters/Demographics/SexAtBirth/SexAtBirth-Ddn",
        "Samples/SampleTumorStatus": "C3DC/Filters/Samples/SampleTumorStatus/SampleTumorStatus_Ddn",
        "Samples/SampleType": "C3DC/Filters/Samples/SampleType/SampleType_Ddn",
        "Diagnosis/Diagnosis": "C3DC/Filters/Diagnosis/Diagnosis/Diagnosis-Ddn",
    },
}

def _wait_c3dc_explore_ready(page):
    """Explore Participants view — left filter panel (Study facet) must be visible."""
    from playwright.sync_api import TimeoutError as PWTimeout

    study_facet = find_test_object("C3DC/Filters/Study/Study_Facet")
    try:
        page.locator(f"xpath={study_facet}").first.wait_for(state="visible", timeout=30000)
        print("[test_runner]   C3DC Explore Participants ready (filter panel visible)")
    except PWTimeout:
        print("[test_runner]   Warning: Explore filter panel not visible after Explore-Tab click")


# C3DC left panel: top-level accordion (Study, Demographics, …) before sub-filter dropdowns.
_C3DC_PANEL_FACETS = {
    "Study/": "C3DC/Filters/Study/Study_Facet",
    "Demographics/": "C3DC/Filters/Demographics/Demographics_Facet",
    "Samples/": "C3DC/Filters/Samples/Samples_Facet",
    "Diagnosis/": "C3DC/Filters/Diagnosis/Diagnosis-Facet",
}


def _ensure_c3dc_panel_expanded(page, obj_name: str, target_xpath: str | None):
    """Open the top-level filter accordion that contains obj_name (Katalon *Facet click)."""
    from playwright.sync_api import TimeoutError as PWTimeout

    if "Filters/" not in obj_name:
        return
    if target_xpath:
        try:
            page.locator(f"xpath={target_xpath}").first.wait_for(state="visible", timeout=1500)
            return
        except PWTimeout:
            pass

    for prefix, facet_repo_key in _C3DC_PANEL_FACETS.items():
        if prefix not in obj_name:
            continue
        try:
            facet_xpath = find_test_object(facet_repo_key)
            floc = page.locator(f"xpath={facet_xpath}").first
            floc.wait_for(state="visible", timeout=8000)
            floc.scroll_into_view_if_needed()
            floc.evaluate("el => el.click()")
            page.wait_for_timeout(600)
            print(f"[test_runner]   Opened C3DC panel: {prefix.rstrip('/')}")
        except Exception as e:
            print(f"[test_runner]   Warning: could not open C3DC panel {prefix}: {e}")
        return


def _ensure_facet_expanded(page, program: str, page_name: str, obj_name: str):
    """
    If obj_name references a filter inside a collapsible facet section that is
    currently collapsed, click the section header to expand it first.
    Silently does nothing if the section is already open or the expander isn't configured.
    """
    from playwright.sync_api import TimeoutError as PWTimeout
    from core.actions import js_click

    effective_obj = obj_name
    if effective_obj.startswith(page_name + "/"):
        effective_obj = effective_obj[len(page_name) + 1:]
    try:
        target_xpath = find_test_object(f"{program}/{page_name}/{effective_obj}")
    except KeyError:
        target_xpath = None

    if program == "C3DC":
        _ensure_c3dc_panel_expanded(page, obj_name, target_xpath)

    expanders = _FACET_EXPANDERS.get(program, {})
    for facet_key, expander_repo_key in expanders.items():
        if facet_key not in obj_name:
            continue
        try:
            expander_xpath = find_test_object(expander_repo_key)
        except KeyError:
            continue
        # Check if any item INSIDE the facet is already visible; if so, it's expanded
        # We probe this by checking whether the facet expander itself is in a closed state.
        # CDS facet expanders are always visible (they're the section headers); the children
        # (Gender-Ddn, etc.) are only visible when the section is open.
        # So: try waiting briefly for the target element; if it's not visible, click expander.
        if target_xpath:
            try:
                loc = page.locator(f"xpath={target_xpath}").first
                loc.wait_for(state="visible", timeout=2000)
                return  # already visible, no need to expand
            except (PWTimeout, Exception):
                pass
        # Section is collapsed — click the expander, then verify the target appeared.
        # Retry once if the first click doesn't stick (React re-render race).
        for attempt in range(2):
            try:
                exp_loc = page.locator(f"xpath={expander_xpath}").first
                exp_loc.wait_for(state="visible", timeout=5000)
                exp_loc.scroll_into_view_if_needed()
                if program == "C3DC":
                    js_click(page, expander_xpath)
                else:
                    exp_loc.click()
                page.wait_for_timeout(800)
                print(f"[test_runner]   Auto-expanded facet: {facet_key} (attempt {attempt+1})")
                # Verify target is now visible; if not, loop and click again
                try:
                    tgt_xpath2 = target_xpath or find_test_object(
                        f"{program}/{page_name}/{effective_obj}"
                    )
                    page.locator(f"xpath={tgt_xpath2}").first.wait_for(state="visible", timeout=2000)
                    return  # target visible — done
                except Exception:
                    pass  # target still not visible; retry click
            except Exception as e:
                print(f"[test_runner]   Warning: could not expand facet '{facet_key}': {e}")
                break
        return


def _verify_static(page, program: str, expected_texts: list) -> bool:
    content = page.content()
    missing = [t for t in expected_texts if t not in content]
    if missing:
        print(f"[test_runner] Static page FAIL — missing: {missing}")
        return False
    print(f"[test_runner] Static page PASS")
    return True
